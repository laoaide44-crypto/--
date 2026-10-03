from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE=Path(__file__).parent
OUT=BASE/'AI软件赛道 - 交付契约台 - EDG - 路演PPT美化版.pptx'
p=Presentation(); p.slide_width= Inches(13.333); p.slide_height= Inches(7.5)
NAVY='132B40'; BLUE='2463A0'; TEAL='087E83'; INK='203345'; MUTED='52687B'; BG='F4F7FA'; WHITE='FFFFFF'; ORANGE='B85C20'
def rect(s,x,y,w,h,c,r=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(c); sh.line.fill.background(); return sh

def text(s,t,x,y,w,h,size=22,c=INK,bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0
 for i,line in enumerate(t.split('\n')):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph(); q.text=line; q.space_after=Pt(10)
  for r in q.runs: r.font.name='Microsoft YaHei'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(c)
 return sh

def page(title,sub):
 s=p.slides.add_slide(p.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 rect(s,0,0,.18,7.5,BLUE); text(s,f'EDG  /  XIHACK 2026  /  {len(p.slides):02}',.65,.3,12,.3,11,BLUE,True)
 text(s,title,.65,.9,12,.7,32,NAVY,True); text(s,sub,.68,1.65,12,.5,17,MUTED)
 text(s,'交付契约台  ·  快递物流主场景',.68,7.02,10,.25,10,MUTED)
 return s

def card(s,x,y,w,title,body,tag=None):
 rect(s,x,y,w,2.75,WHITE,True); rect(s,x+.2,y+.25,.05,.55,TEAL)
 text(s,title,x+.4,y+.25,w-.7,.65,23,NAVY,True)
 text(s,body,x+.3,y+1,w-.6,1.5,18,INK)
 if tag: text(s,tag,x+.3,y+2.35,w-.6,.3,12,TEAL,True)

def note(s,t): s.notes_slide.notes_text_frame.text=t

s=p.slides.add_slide(p.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY)
rect(s,9.7,0,3.63,7.5,BLUE); text(s,'EDG  /  一人团队',.8,.65,8,.5,18,'B8D6EF',True)
text(s,'交付契约台',.8,2,8.5,1,48,WHITE,True); text(s,'AI 做完的活，\n如何被验收？',.85,3.3,8,1.6,30,WHITE)
text(s,'AI软件赛道  ·  XiHack 2026',.85,6.55,8,.4,15,'B8D6EF')
for y,t in [(1.8,'有依据'),(3.3,'可复算'),(4.8,'能撤回')]: text(s,t,10.15,y,2.8,.8,25,WHITE,True)
note(s,'我一个人组成 EDG 队。作品叫交付契约台，关注 AI 的结果如何被验收。正式演示以快递物流为主场景。建议 20 秒。')

s=page('答案很快，验收却卡在三个问题','企业材料分散；口径、数字和结论需要逐层核对。')
card(s,.7,2.55,3.85,'依据在哪？','群聊、表格与工单\n结论要回到原句与原行。','01  来源可追溯')
card(s,4.75,2.55,3.85,'数字怎么算？','口径不同不能硬比\n计算链要能人工复算。','02  算术可检查')
card(s,8.8,2.55,3.85,'能承诺什么？','缺少证据就停在方案设计\n不把假设当作客户成果。','03  交付有边界')
text(s,'先把证据说清楚，再决定下一步。',.8,6.1,11.8,.6,26,BLUE,True)
note(s,'企业给来的材料往往很乱。回答得像不像还不够，我们需要知道依据在哪、数字怎么算、目前能承诺什么。建议 35 秒。')

s=page('一条流程，把材料变成可核对的交付','赛前已有底座：事实抽取、来源回指、行业测算、证据门槛与契约草稿。')
for i,(a,b) in enumerate([('原始材料','群聊 / 工单 / 表格'),('事实与缺口','出处 / 来源 / 追问'),('确定性测算','公开基准 / 参数'),('交付契约','范围 / 验收 / 限制')]):
 x=.75+i*3.15; rect(s,x,2.65,2.8,2.4,WHITE,True); text(s,f'0{i+1}',x+.25,2.9,2,.4,18,TEAL,True); text(s,a,x+.25,3.5,2.4,.55,24,NAVY,True); text(s,b,x+.25,4.2,2.4,.55,15,MUTED)
 if i<3: text(s,'→',x+2.85,3.45,.3,.5,22,BLUE,True)
text(s,'模型：抽取与表达     |     代码：校验、算术与规则',.8,5.9,12,.6,23,BLUE,True)
note(s,'这条底座属于赛前已有成果，我会单列申报。模型负责抽取与表达，确定性代码负责算术和规则。建议 35 秒。')

s=page('主场景：快递物流单票成本诊断','把成本定义、公开参照和计算步骤同时交给用户。')
rect(s,.75,2.5,7.4,3.55,NAVY,True); text(s,'单票差额 × 年件量',1.1,3.05,6.7,.7,32,WHITE,True)
text(s,'口径先对齐\n算术可展开\n来源可回指',1.1,4.05,6.7,1.55,23,'CFE1EF')
text(s,'必须保留的边界',8.65,2.7,3.9,.5,23,BLUE,True)
text(s,'口径未对齐\n→ 只到方案设计\n\n假设优化空间\n≠ 已验证节省',8.65,3.5,3.85,2.7,22,INK)
note(s,'演示先选快递物流，展开基准与算术。录屏数值以实际画面为准，不读脚本中的固定收益数字。建议 40 秒。')

s=page('现场新增：让 CSV 约束计算结果','上传用途显式分流；合成样本验证，不代表客户真实业务。')
rect(s,.75,2.55,4,3.6,BLUE,True); text(s,'30.00',1.1,3,3.3,1,48,WHITE,True); text(s,'小时 / 文件内已解析记录',1.1,4.15,3.35,.65,19,WHITE)
text(s,'1800 分钟\n净工时 = 结束 − 开始 − 休息',1.1,5,3.4,.85,16,'D6E8F8')
text(s,'合成测试材料 · 非真实客户数据',5.25,2.7,7,.5,22,NAVY,True)
text(s,'员工小计：14 / 4 / 8 / 4 小时\n门店小计：18 / 12 小时\n每个数字回指到行号与字段',5.25,3.4,7,1.9,23,INK)
text(s,'计划排班 ≠ 实际出勤   ·   合计 ≠ 范围完整',5.25,5.65,7,.7,18,ORANGE,True)
note(s,'现场新增 CSV 路径与原有文本路径并存。上传有效样本得到 30 小时，展示行字段回指。它只证明文件内计划工时，不证明考勤或真实性。建议 45 秒。')

s=page('异常出现时，系统必须收住结论','完全重复去重计一次；非完全重复的重叠不靠区间并集掩盖。')
card(s,.7,2.6,3.85,'字段错误','E1–E5 混合样本\n11.00 小时无异常小计','阻断应用')
card(s,4.75,2.6,3.85,'跨门店重叠','W2 · 未解决冲突\n7.00 小时无异常小计','阻断应用')
card(s,8.8,2.6,3.85,'完全重复','W1 · 重复只计一次\n11.00 小时','保留原始行号映射')
text(s,'少给一个夸大结论，比多给一个漂亮数字更有价值。',.8,6.15,11.8,.6,24,BLUE,True)
note(s,'这页数值是合成样本的确定性结果。错误与重叠样本不能确认应用。演示任选重叠样本停留，指出冲突行与确认按钮状态。建议 40 秒。')

s=page('采用一个版本，也能解释并撤回它','确认采用 ≠ 真实性已验证。证据版本独立保存来源、合成标识与状态。')
for x,a,b in [(1,'v1','1800 分钟 / 30.00 小时'),(5,'v2','1830 分钟 / 30.50 小时'),(9,'撤回 → v1','保留 v2 历史')]:
 rect(s,x,2.7,3.2,2.3,WHITE,True); text(s,a,x+.25,3.1,2.7,.65,30,BLUE,True); text(s,b,x+.25,4.15,2.7,.6,17,INK)
text(s,'→',4.35,3.45,.5,.7,25,TEAL,True); text(s,'→',8.35,3.45,.5,.7,25,TEAL,True)
text(s,'哈希 + 规则版本 + 覆盖范围 + 变化依据',1,5.55,11.5,.5,23,NAVY,True)
text(s,'页面与导出来自同一个数据载荷',1,6.25,11.5,.4,19,TEAL,True)
note(s,'v2 示例仅把原始合成样本第一行休息从 60 分改成 30 分，工时增加 30 分钟。这是测试变化，不是优化成果。撤回后看当前版本 JSON 与导出，不把上传预览当成当前版本。建议 45 秒。')

s=page('可复核的工程结果，明确的贡献边界','一人完成项目组织与交付；AI 辅助开发已通过真实提交和测试留痕。')
card(s,.75,2.6,5.95,'赛前已有','事实抽取 / 回指 / 行业测算\n证据门槛 / 契约 / 离线缓存','PRE_EXISTING 声明')
card(s,6.9,2.6,5.65,'现场新增','CSV 校验与计算\n来源回指 / 证据版本与撤回','73925f7 · 5c76aa6 · b5f3fec')
text(s,'三段直接测试通过  ·  UI 源码 smoke 通过  ·  编译通过',.9,5.9,11.8,.5,22,BLUE,True)
text(s,'合成测试验证计算行为；尚无真实客户试点成效证明。',.9,6.55,11.8,.35,16,MUTED)
note(s,'不要称所有功能为比赛现场完成。三段验收为直接 Python 测试，UI smoke 为源码检查，不能当成完整浏览器自动化验证。建议 35 秒。')

s=page('让结果有依据，让交付有边界','EDG · 一人团队 · 交付契约台')
text(s,'有依据   /   可复算   /   能撤回',.85,2.8,12,1,37,BLUE,True)
text(s,'物流作为主场景，餐饮作为对照\n下一步：在授权、脱敏的数据上对齐口径，验证真实场景。',.9,4.2,11.6,1.4,24,INK)
text(s,'AI软件赛道 - 交付契约台 - EDG',.9,6.15,11.5,.5,19,MUTED)
note(s,'结束前可短暂切到餐饮展示行业差异。不要说已经证明真实降本。建议 20 秒。')
p.save(OUT)
print('created',len(p.slides),'slides')
