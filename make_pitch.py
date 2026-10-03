from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

OUT = Path(__file__).parent / 'AI软件赛道 - 交付契约台 - EDG - 路演PPT.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
NAVY = RGBColor(21, 47, 74)
BLUE = RGBColor(31, 78, 121)
INK = RGBColor(32, 42, 52)
MUTED = RGBColor(91, 105, 117)
BG = RGBColor(247, 249, 251)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(211, 126, 45)


def textbox(slide, text, x, y, w, h, size=20, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None: p.alignment = align
    for i, line in enumerate(text.split('\n')):
        if i: p = tf.add_paragraph()
        p.text = line; p.space_after = Pt(8)
        for r in p.runs:
            r.font.name = 'Aptos'; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box


def slide(title, kicker=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = BG
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.18))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    if kicker: textbox(s, kicker.upper(), .7, .48, 11.8, .3, 10, ACCENT, True)
    textbox(s, title, .7, .8, 11.8, .65, 30, NAVY, True)
    textbox(s, 'AI软件赛道 · 交付契约台 · EDG', .7, 7.08, 6, .2, 9, MUTED)
    return s


def bullets(s, items, x=.9, y=1.8, w=11.5, size=20):
    text = '\n'.join('• ' + i for i in items)
    textbox(s, text, x, y, w, 4.8, size, INK)

s=prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
textbox(s, 'AI软件赛道', .9, 1.15, 5, .4, 16, RGBColor(160, 197, 224), True)
textbox(s, '交付契约台', .9, 1.8, 11.5, 1.0, 42, WHITE, True)
textbox(s, 'AI 干完的活，如何被验收', .95, 3.1, 10, .55, 27, RGBColor(224, 232, 238))
textbox(s, '快递物流主场景 · 餐饮门店对照 · EDG', .95, 5.8, 10, .4, 16, RGBColor(190, 207, 219))

s=slide('问题：AI 做完的活如何验收', '01 · 场景')
bullets(s, ['企业材料杂乱，事实、口径和结论容易混在一起。','快递物流中，单票成本差异乘上件量，会形成很大的年度影响。','先给可复算结果，再让客户决定是否提供更多数据。'], size=23)

s=slide('产品：从材料到可签字的交付契约', '02 · 产品')
bullets(s, ['抽取带出处、带来源等级的事实。','只追问材料还没有回答的问题。','公开基准 + 少量参数完成确定性测算。','生成写清证据、边界和验收条件的交付契约。','餐饮门店作为第二行业，验证换行业不换界面与契约骨架。'], size=21)

s=slide('物流主场景：先算清楚，再谈优化', '03 · 主场景')
textbox(s, '快递物流 · 单票成本诊断', .9, 1.75, 11, .45, 25, BLUE, True)
bullets(s, ['公开基准、逐步算术、来源回指全部可展开。','口径未对齐时只到“方案设计”，不承诺收益。','不把假设情景说成已验证客户结果。'], y=2.55, size=22)

s=slide('现场新增：结构化排班 CSV 能力', '04 · 现场新增')
bullets(s, ['显式选择“作为文本分析 / 作为排班表校验”。','UTF-8/BOM CSV 的表头、字段、记录、重复和重叠校验。','有效行净分钟/小时、员工/门店小计、来源行和字段回指。','阻断或未解决重叠时只给“无异常记录小计”，不替换完整基线。','合成样本明确标识为测试材料，不代表真实客户证据。'], size=19)

s=slide('证据版本：确认才应用，变化可追溯', '05 · 证据闭环')
bullets(s, ['用户显式确认后创建 vN 证据版本。','记录文件哈希、规则版本、声明/实际范围、来源、合成标识和计算结果。','支持变化报告、版本切换、撤回和同源导出。','页面与导出逐条一致；导入不会自动升级真实性或证据门槛。'], size=20)

s=slide('技术与验证：确定性代码守住边界', '06 · 工程')
bullets(s, ['通用内核：事实、证据、门槛、契约；行业模块：物流与餐饮分别计算。','确定性代码负责解析、校验、算术和版本；模型不做算术、不决定证据升级。','现场提交：73925f7 · 5c76aa6 · b5f3fec；贡献记录：8dc2bce；提交准备：aae5242。','Stage 1/2/3、UI smoke 和编译验证均通过。'], size=19)

s=slide('现场演示与提交', '07 · DEMO')
bullets(s, ['打开 http://127.0.0.1:8502，选择快递物流。','上传 SYNTHETIC_VALID_two_stores.csv，展示 30.00 小时和来源回指。','切换错误/重叠样本，展示阻断与“不应用”边界。','显式确认 v1，修改后生成 v2，再撤回并导出。','结尾切换餐饮作为对照；演示材料明确为合成/非真实客户数据。'], size=19)
textbox(s, '作品提交标识：AI软件赛道 - 交付契约台 - EDG', .9, 6.25, 11.5, .45, 19, BLUE, True)

prs.save(OUT)
print(OUT)
